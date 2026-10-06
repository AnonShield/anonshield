"""
Regression tests for flags and modes that used to leave PII in the output or
ignore the user's selection. Most run in-process with the regex strategy (no
model download); the CLI cases also use --anonymization-strategy regex.
"""
import csv
import json
import re
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from src.anon.cache_manager import CacheManager
from src.anon.config import Global
from src.anon.engine import AnonymizationOrchestrator, chunked_ner, load_custom_recognizers
from src.anon.entity_detector import EntityDetector
from src.anon.entity_selection import resolve_entity_selection
from src.anon.hash_generator import HashGenerator
from src.anon.processors import ProcessorRegistry

PROJECT_ROOT = Path(__file__).parent.parent
KEY = "test-key-12345678901234567890123456789012"
EMAIL, EMAIL2, IP = "maria.oliveira@example.com", "john.smith@acme-corp.com", "192.168.10.45"


def regex_orchestrator(preserve=(), allow=(), selected=None, slug_length=8, secret_key=KEY):
    return AnonymizationOrchestrator(
        lang="en", db_context=None, allow_list=list(allow),
        entities_to_preserve=list(Global.NON_PII_ENTITIES | set(preserve)),
        slug_length=slug_length, strategy_name="regex",
        cache_manager=CacheManager(use_cache=True, max_cache_size=1000),
        hash_generator=HashGenerator(secret_key=secret_key),
        entities_to_anonymize=selected,
    )


def process(path, orchestrator, out_dir, **kwargs):
    processor = ProcessorRegistry.get_processor(str(path), orchestrator, output_dir=str(out_dir),
                                                overwrite=True, **kwargs)
    return Path(processor.process())


# ---------------------------------------------------------------------------
# Entity selection and detection policy
# ---------------------------------------------------------------------------

def test_entities_selection_is_positive():
    sel = resolve_entity_selection({"EMAIL_ADDRESS", "IP_ADDRESS", "MONEY"}, entities=["email_address", "MONEY", "NOPE"])
    assert sel.entities_to_anonymize == {"EMAIL_ADDRESS", "MONEY"}
    assert "MONEY" not in sel.entities_to_preserve  # selectable despite being non-PII by default
    assert sel.unknown == ["NOPE"]
    assert resolve_entity_selection({"EMAIL_ADDRESS"}, entities=[]).entities_to_anonymize == set()


def test_unselected_label_outside_supported_list_is_not_anonymized():
    detector = EntityDetector([], set(), set(), entities_to_anonymize={"EMAIL_ADDRESS"})
    text = "On 2024-01-01 write to a@b.com"
    ents = [{"start": 3, "end": 13, "label": "DATE", "text": "2024-01-01", "score": 0.99},
            {"start": 23, "end": 30, "label": "EMAIL_ADDRESS", "text": "a@b.com", "score": 1.0}]
    assert [e["label"] for e in detector.finalize(text, ents)] == ["EMAIL_ADDRESS"]


def test_preserved_type_keeps_its_span():
    # A preserved IP must not come back under another label (Presidio's phone recognizer).
    detector = EntityDetector([], {"IP_ADDRESS"}, set())
    ents = [{"start": 0, "end": 13, "label": "IP_ADDRESS", "text": IP, "score": 0.85},
            {"start": 0, "end": 13, "label": "PHONE_NUMBER", "text": IP, "score": 0.4}]
    assert detector.finalize(IP, ents) == []


def test_allow_list_covers_matches_inside_the_term():
    detector = EntityDetector([], set(), {EMAIL})
    text = f"mail {EMAIL} now"
    start = text.index(EMAIL)
    ents = [{"start": start, "end": start + len(EMAIL), "label": "EMAIL_ADDRESS", "text": EMAIL, "score": 1.0},
            {"start": start, "end": start + 14, "label": "HOSTNAME", "text": "maria.oliveira", "score": 0.6}]
    assert detector.finalize(text, ents) == []


def test_partial_overlap_is_covered_not_dropped():
    detector = EntityDetector([], set(), set())
    text = "Projeto Fenix rocks"
    ents = [{"start": 0, "end": 7, "label": "ORGANIZATION", "text": "Projeto", "score": 0.9},
            {"start": 5, "end": 13, "label": "PROJECT", "text": "to Fenix", "score": 0.8}]
    merged = detector.finalize(text, ents)
    assert len(merged) == 1 and merged[0]["text"] == "Projeto Fenix"


def test_regex_priority_scores_stay_valid():
    recognizers = load_custom_recognizers(["en", "pt"], regex_priority=True)
    assert all(p.score <= 1.0 for r in recognizers for p in r.patterns)


def test_hash_generator_slug_zero_needs_no_key_and_explicit_key_is_used():
    assert HashGenerator(secret_key=None).generate_slug("x", 0) == ("", "")
    a = HashGenerator(secret_key="key-a").generate_slug("john", 8)[0]
    b = HashGenerator(secret_key="key-b").generate_slug("john", 8)[0]
    assert a != b


def test_chunked_ner_sees_the_whole_text_and_keeps_names_cut_by_a_window():
    class Tok:
        def __call__(self, text, **_):
            return {"offset_mapping": [(m.start(), m.end()) for m in re.finditer(r"\S+", text)]}

    class Pipe:
        tokenizer = Tok()

        def __call__(self, text):
            return [{"entity_group": "PER", "start": m.start(), "end": m.end(), "score": 0.99, "word": m.group()}
                    for m in re.finditer(r"Maria Oliveira", text)]

    text = " ".join(["filler"] * 395) + " Maria Oliveira " + " ".join(["filler"] * 600) + " Maria Oliveira"
    found = [text[e["start"]:e["end"]] for e in chunked_ner(Pipe(), text, chunk_tokens=400, overlap=50)]
    assert found == ["Maria Oliveira", "Maria Oliveira"]


def test_cuda_arch_support_matches_the_build():
    from src.anon.device import _arch_supported
    cu126 = ["sm_50", "sm_60", "sm_70", "sm_75", "sm_80", "sm_86", "sm_90"]
    cu130 = ["sm_75", "sm_80", "sm_86", "sm_90", "sm_100", "sm_120", "compute_120"]
    assert not _arch_supported(12, 0, cu126)   # RTX 50xx on a CUDA 12.6 build
    assert _arch_supported(12, 0, cu130)
    assert _arch_supported(8, 9, cu130)        # RTX 40xx runs the sm_86 cubin
    assert not _arch_supported(6, 1, cu130)    # GTX 10xx on a CUDA 13 build
    assert _arch_supported(6, 1, cu126)


# ---------------------------------------------------------------------------
# File processors
# ---------------------------------------------------------------------------

def test_csv_preserve_row_context_anonymizes(tmp_path):
    src = tmp_path / "t.csv"
    src.write_text(f"id,email\n1,{EMAIL}\n2,{EMAIL2}\n", encoding="utf-8")
    out = process(src, regex_orchestrator(), tmp_path / "out", preserve_row_context=True)
    content = out.read_text(encoding="utf-8")
    assert EMAIL not in content and EMAIL2 not in content


@pytest.mark.parametrize("prc", [False, True])
def test_csv_excluded_column_stays_intact(tmp_path, prc):
    src = tmp_path / "t.csv"
    src.write_text(f"contact,audit\n{EMAIL},{EMAIL}\n{EMAIL2},{EMAIL2}\n", encoding="utf-8")
    out = process(src, regex_orchestrator(), tmp_path / "out", preserve_row_context=prc,
                  anonymization_config={"fields_to_exclude": ["audit"]})
    rows = list(csv.DictReader(out.open(encoding="utf-8")))
    assert [r["audit"] for r in rows] == [EMAIL, EMAIL2]
    assert all(r["contact"].startswith("[EMAIL_ADDRESS_") for r in rows)


def test_xlsx_excluded_column_stays_intact(tmp_path):
    src = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in (["contact", "audit"], [EMAIL, EMAIL]):
        ws.append(row)
    wb.save(src)
    out = process(src, regex_orchestrator(), tmp_path / "out", anonymization_config={"fields_to_exclude": ["Sheet1.B"]})
    ws = openpyxl.load_workbook(out).active
    assert ws["B2"].value == EMAIL
    assert ws["A2"].value.startswith("[EMAIL_ADDRESS_")


def test_txt_context_mode_decides_per_line(tmp_path):
    # The first line of the batch is a stoplist word; it used to decide for the whole batch.
    src = tmp_path / "t.txt"
    src.write_bytes(f"high\r\n\r\nwrite to {EMAIL}\r\n".encode())
    out = process(src, regex_orchestrator(), tmp_path / "out", preserve_row_context=True, min_word_length=3)
    content = out.read_bytes().decode()
    assert EMAIL not in content
    assert content.startswith("high\r\n\r\n") and content.endswith("]\r\n")


def test_structured_config_does_not_disable_text_files(tmp_path):
    src = tmp_path / "t.txt"
    src.write_text(f"write to {EMAIL}\n", encoding="utf-8")
    out = process(src, regex_orchestrator(), tmp_path / "out",
                  anonymization_config={"fields_to_anonymize": ["description"]})
    assert EMAIL not in out.read_text(encoding="utf-8")


def test_xml_paths_are_root_first_and_comments_are_handled(tmp_path):
    src = tmp_path / "t.xml"
    src.write_text(f'<tickets><ticket reporter="{EMAIL}"><audit>{EMAIL}</audit>'
                   f'<!-- escalated by {EMAIL2} --></ticket></tickets>', encoding="utf-8")
    out = process(src, regex_orchestrator(), tmp_path / "out",
                  anonymization_config={"fields_to_exclude": ["tickets/ticket/audit"]})
    content = out.read_text(encoding="utf-8")
    assert f"<audit>{EMAIL}</audit>" in content
    assert f'reporter="{EMAIL}"' not in content
    assert EMAIL2 not in content


def test_json_forced_number_is_anonymized(tmp_path):
    src = tmp_path / "t.json"
    src.write_text(json.dumps({"user": {"cpf": 12345678909, "age": 40}}), encoding="utf-8")
    out = process(src, regex_orchestrator(), tmp_path / "out",
                  anonymization_config={"force_anonymize": {"user.cpf": {"entity_type": "BR_CPF"}}})
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["user"]["cpf"].startswith("[BR_CPF_")
    assert data["user"]["age"] == 40


def test_json_listed_number_without_pii_keeps_its_type(tmp_path):
    src = tmp_path / "t.json"
    src.write_text(json.dumps({"user": {"age": 40}}), encoding="utf-8")
    out = process(src, regex_orchestrator(), tmp_path / "out",
                  anonymization_config={"fields_to_anonymize": ["user.age"]})
    assert json.loads(out.read_text(encoding="utf-8"))["user"]["age"] == 40


def test_invalid_json_is_an_error(tmp_path):
    src = tmp_path / "t.json"
    src.write_text("{not json", encoding="utf-8")
    with pytest.raises(ValueError):
        process(src, regex_orchestrator(), tmp_path / "out")


def test_forced_type_list_works_without_presidio(tmp_path):
    src = tmp_path / "t.json"
    src.write_text(json.dumps({"host": "10.0.0.7"}), encoding="utf-8")
    cfg = {"force_anonymize": {"host": {"entity_type": ["HOSTNAME", "IP_ADDRESS"]}}}
    out = process(src, regex_orchestrator(), tmp_path / "out", anonymization_config=cfg)
    assert json.loads(out.read_text(encoding="utf-8"))["host"].startswith("[IP_ADDRESS_")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def run_cli(tmp_path, text, *args, env_key=True):
    src = tmp_path / "in.txt"
    src.write_text(text, encoding="utf-8")
    out_dir = tmp_path / "out"
    env = {k: v for k, v in __import__("os").environ.items() if not k.startswith("ANON_SECRET_KEY")}
    if env_key:
        env["ANON_SECRET_KEY"] = KEY
    cmd = [sys.executable, str(PROJECT_ROOT / "anon.py"), str(src), "--anonymization-strategy", "regex",
           "--db-mode", "in-memory", "--no-report", "--overwrite", "--output-dir", str(out_dir), *args]
    result = subprocess.run(cmd, capture_output=True, text=True, cwd=str(PROJECT_ROOT), env=env, timeout=120)
    out = out_dir / "anon_in.txt"
    return result, out.read_text(encoding="utf-8") if out.exists() else ""


def test_cli_slug_zero_without_key(tmp_path):
    result, content = run_cli(tmp_path, f"write to {EMAIL}\n", "--slug-length", "0", env_key=False)
    assert result.returncode == 0, result.stdout + result.stderr
    assert content == "write to [EMAIL_ADDRESS]\n"


def test_cli_word_list_and_custom_patterns_with_selection(tmp_path):
    words = tmp_path / "words.json"
    words.write_text(json.dumps({"PROJECT": ["Projeto Fenix"]}), encoding="utf-8")
    patterns = tmp_path / "patterns.yaml"
    patterns.write_text("- entity_type: TICKET_ID\n  pattern: 'TICKET-\\d{6}'\n  score: 0.95\n", encoding="utf-8")
    text = f"Projeto Fenix TICKET-482913 {EMAIL}\n"
    result, content = run_cli(tmp_path, text, "--word-list", str(words), "--custom-patterns", str(patterns),
                              "--entities", "PROJECT,TICKET_ID")
    assert result.returncode == 0, result.stdout + result.stderr
    assert "Projeto Fenix" not in content and "TICKET-482913" not in content
    assert EMAIL in content  # not selected


def test_cli_lang_pt_entities_only_email(tmp_path):
    result, content = run_cli(tmp_path, f"CPF 123.456.789-09 {EMAIL}\n", "--lang", "pt", "--entities", "EMAIL_ADDRESS")
    assert result.returncode == 0, result.stdout + result.stderr
    assert "123.456.789-09" in content and EMAIL not in content



def anonymize_text(tmp_path, text, **kwargs):
    """Text through the regex strategy with type-only labels (--slug-length 0)."""
    src = tmp_path / "in.txt"
    src.write_text(text, encoding="utf-8")
    orchestrator = regex_orchestrator(slug_length=0, secret_key=None, **kwargs)
    return process(src, orchestrator, tmp_path / "out").read_text(encoding="utf-8")


@pytest.mark.parametrize("phone", ["+1 212 555 0198", "+44 20 7946 0958", "(212) 555-0198", "212-555-0198",
                                   "+55 51 99999-9999", "(51) 99999-9999"])
def test_regex_strategy_catches_international_and_us_phones(tmp_path, phone):
    assert anonymize_text(tmp_path, f"Call {phone} today\n") == "Call [PHONE_NUMBER] today\n"


@pytest.mark.parametrize("text", ["version 2.10.3 build 2024-10-06 port 8080", "ticket 12345 at 10:30, 3 hosts"])
def test_phone_patterns_leave_plain_numbers_alone(tmp_path, text):
    assert anonymize_text(tmp_path, text + "\n", selected={"PHONE_NUMBER"}) == text + "\n"


def test_output_name_keeps_spaces_and_inner_dots(tmp_path):
    from src.anon.processors import get_output_path
    assert Path(get_output_path("/data/nota final.v2.txt", ".txt", output_dir=str(tmp_path))).name == "anon_nota final.v2.txt"
    assert Path(get_output_path("../../etc/passwd", ".txt", output_dir=str(tmp_path))).parent == tmp_path.resolve()
    with pytest.raises(ValueError):
        get_output_path("..", ".txt", output_dir=str(tmp_path))
