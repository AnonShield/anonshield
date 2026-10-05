"""
End-to-end runs of the CLI over a directory with every supported text format
(txt, csv, json, jsonl, xml, xlsx, docx, pdf), for each strategy and the main
flag combinations. Each run is checked for PII left in the output, which is how
most of the regressions fixed alongside these tests showed up: the run
"succeeded" and the file came out unchanged.

Images are left out (OCR needs Tesseract, see the "ocr" marker).
"""
import csv
import json
import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pymupdf as fitz
import pytest
from docx import Document

PROJECT_ROOT = Path(__file__).parent.parent
KEY = "test-key-12345678901234567890123456789012"

EMAILS = ["maria.oliveira@example.com", "john.smith@acme-corp.com"]
IPS = ["192.168.10.45", "10.20.30.40"]
NAMES = ["Maria Oliveira", "John Smith"]
LINES = [
    f"Incident reported by {NAMES[0]} ({EMAILS[0]}) from host {IPS[0]}.",
    f"{NAMES[1]} <{EMAILS[1]}> confirmed the scan from {IPS[1]}.",
]
OUTPUTS = {
    "notes.txt": "anon_notes.txt", "tickets.csv": "anon_tickets.csv", "tickets.json": "anon_tickets.json",
    "events.jsonl": "anon_events.jsonl", "tickets.xml": "anon_tickets.xml", "tickets.xlsx": "anon_tickets.xlsx",
    "report.docx": "anon_report.txt",
}


@pytest.fixture(scope="module")
def data_dir(tmp_path_factory):
    d = tmp_path_factory.mktemp("e2e_data")
    # A stoplist word and a blank line first: they used to decide for the whole batch.
    (d / "notes.txt").write_text("high\n\n" + "\n".join(LINES) + "\n", encoding="utf-8")
    rows = [["id", "category", "contact", "audit", "description"],
            ["1", "CAT5", EMAILS[0], EMAILS[0], LINES[0]],
            ["2", "CAT3", EMAILS[1], "auditor@example.org", LINES[1]]]
    with open(d / "tickets.csv", "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)
    tickets = [{"id": i, "reporter": {"name": NAMES[i], "email": EMAILS[i]}, "ip": IPS[i], "notes": LINES[i]}
               for i in range(2)]
    (d / "tickets.json").write_text(json.dumps({"tickets": tickets}), encoding="utf-8")
    (d / "events.jsonl").write_text("".join(json.dumps(t) + "\n" for t in tickets), encoding="utf-8")
    (d / "tickets.xml").write_text(
        "<tickets>" + "".join(f'<ticket reporter="{EMAILS[i]}"><notes>{LINES[i]}</notes></ticket>' for i in range(2))
        + "</tickets>", encoding="utf-8")
    wb = openpyxl.Workbook()
    for row in rows:
        wb.active.append(row)
    wb.save(d / "tickets.xlsx")
    doc = Document()
    doc.add_paragraph("description")
    for line in LINES:
        doc.add_paragraph(line)
    doc.save(d / "report.docx")
    pdf_dir = tmp_path_factory.mktemp("e2e_pdf")
    pdf = fitz.open()
    page = pdf.new_page()
    for i, line in enumerate(["severity"] + LINES):
        page.insert_text((72, 72 + 30 * i), line, fontsize=10)
    pdf.save(pdf_dir / "report.pdf")
    return d, pdf_dir


def run_dir(input_dir: Path, out_dir: Path, *args: str) -> subprocess.CompletedProcess:
    env = {**os.environ, "ANON_SECRET_KEY": KEY}
    cmd = [sys.executable, str(PROJECT_ROOT / "anon.py"), str(input_dir), "--output-dir", str(out_dir),
           "--db-mode", "in-memory", "--no-report", "--overwrite", *args]
    return subprocess.run(cmd, capture_output=True, text=True, cwd=str(PROJECT_ROOT), env=env, timeout=900)


def read_output(path: Path) -> str:
    if path.suffix == ".xlsx":
        wb = openpyxl.load_workbook(path)
        return "\n".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    return path.read_text(encoding="utf-8")


def run_all(data_dir, tmp_path, *args):
    """Run the main directory and the PDF directory (docx and pdf share an output name)."""
    main_dir, pdf_dir = data_dir
    outputs = {}
    for name, (src, sub) in {"main": (main_dir, "main"), "pdf": (pdf_dir, "pdf")}.items():
        result = run_dir(src, tmp_path / sub, *args)
        log = result.stdout + result.stderr
        assert result.returncode == 0, log[-3000:]
        assert " - ERROR - " not in log and "Traceback" not in log, log[-3000:]
    for src_name, out_name in OUTPUTS.items():
        p = tmp_path / "main" / out_name
        assert p.exists(), f"no output for {src_name}"
        outputs[src_name] = read_output(p)
    outputs["report.pdf"] = read_output(tmp_path / "pdf" / "anon_report.txt")
    return outputs


def assert_no_leak(outputs, needles):
    leaks = {f: [n for n in needles if n in text] for f, text in outputs.items()}
    assert not any(leaks.values()), {f: l for f, l in leaks.items() if l}


STRATEGIES = ["filtered", "presidio", "hybrid", "standalone", "regex"]


@pytest.mark.parametrize("strategy", STRATEGIES)
@pytest.mark.parametrize("mode", [[], ["--preserve-row-context"]], ids=["dedup", "row-context"])
def test_no_pii_left(data_dir, tmp_path, strategy, mode):
    outputs = run_all(data_dir, tmp_path, "--anonymization-strategy", strategy, *mode)
    assert_no_leak(outputs, EMAILS + IPS + ([] if strategy == "regex" else NAMES))


@pytest.mark.parametrize("strategy", ["filtered", "standalone"])
def test_allow_list_and_preserve(data_dir, tmp_path, strategy):
    outputs = run_all(data_dir, tmp_path, "--anonymization-strategy", strategy,
                      "--allow-list", EMAILS[0], "--preserve-entities", "IP_ADDRESS")
    for f, text in outputs.items():
        assert EMAILS[0] in text, f"{f}: allow-listed e-mail was changed"
        assert IPS[0] in text and "[IP_ADDRESS" not in text, f"{f}: preserved IP was changed"
    assert_no_leak(outputs, [EMAILS[1]] + NAMES)


@pytest.mark.parametrize("strategy", ["filtered", "standalone"])
def test_entities_selection(data_dir, tmp_path, strategy):
    outputs = run_all(data_dir, tmp_path, "--anonymization-strategy", strategy, "--entities", "EMAIL_ADDRESS")
    for f, text in outputs.items():
        assert IPS[0] in text and NAMES[0] in text, f"{f}: a type outside --entities was changed"
    assert_no_leak(outputs, EMAILS)


def test_structured_exclusion(data_dir, tmp_path):
    cfg = tmp_path / "cfg.json"
    cfg.write_text(json.dumps({"fields_to_exclude": ["audit", "category"]}), encoding="utf-8")
    run_all(data_dir, tmp_path, "--anonymization-config", str(cfg))
    rows = list(csv.DictReader(open(tmp_path / "main" / "anon_tickets.csv", encoding="utf-8")))
    assert [r["audit"] for r in rows] == [EMAILS[0], "auditor@example.org"]
    assert [r["category"] for r in rows] == ["CAT5", "CAT3"]
    assert all(r["contact"].startswith("[EMAIL_ADDRESS_") for r in rows)
